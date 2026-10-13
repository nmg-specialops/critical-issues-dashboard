"""Header-based partner issue reading and preservation-aware workbook updates."""
from io import BytesIO
from datetime import datetime, date, timezone
from copy import copy
from openpyxl import load_workbook

PARTNERS={'SPCL':'Serendipalm','SPOL':'Serendipol','PM':'Pavitramenthe'}
FIELDS=['Project','Topic','Root cause','Problem','Effects','Solutions','Responsible','Date Identified','Targeted Resolution','Notes','Status','Update']
DATES={'Date Identified','Targeted Resolution'}
HISTORY=['Partner','Project','Issue','Timestamp UTC','User','Field','Previous value','New value','Note']
ALIASES={'project':'Project','topic':'Topic','root cause':'Root cause','problem':'Problem','effects':'Effects','solution':'Solutions','solutions':'Solutions','responsible':'Responsible','date identified':'Date Identified','target resolution':'Targeted Resolution','targeted resolution':'Targeted Resolution','notes':'Notes','pertinent notes':'Notes','status':'Status','update':'Update','updates':'Update'}

def text(v):
    if v is None:return ''
    if isinstance(v,(datetime,date)):return v.strftime('%Y-%m-%d')
    return str(v).strip()

def anchor(ws,r,c):
    for m in ws.merged_cells.ranges:
        if m.min_row<=r<=m.max_row and m.min_col<=c<=m.max_col:return m.min_row,m.min_col
    return r,c

def mapping(ws):
    cols={}
    for cell in ws[4]:
        field=ALIASES.get(text(cell.value).lower())
        if field:
            if field in cols:raise ValueError(f'{ws.title}: duplicate column for {field}.')
            cols[field]=cell.column
    if not {'Project','Topic','Problem'}<=cols.keys():raise ValueError(f'{ws.title}: Project, Topic and Problem headers are required in row 4.')
    return cols

def read(data):
    wb=load_workbook(BytesIO(data));records=[]
    for partner in PARTNERS:
        if partner not in wb:raise ValueError(f'Missing sheet: {partner}')
        ws=wb[partner];cols=mapping(ws);groups={}
        for r in range(5,ws.max_row+1):
            if not any(ws.cell(r,c).value is not None for c in cols.values()):continue
            start,_=anchor(ws,r,cols['Problem'])
            if not text(ws.cell(start,cols['Problem']).value):
                raise ValueError(f'{partner} row {r}: enter a Problem for this issue.')
            groups.setdefault(start,[]).append(r)
        for start,rows in groups.items():
            for m in ws.merged_cells.ranges:
                if m.min_row==start and m.min_col<=cols['Problem']<=m.max_col:
                    rows=list(range(start,max(max(rows),m.max_row)+1))
            rec={f:'' for f in FIELDS}
            for field,col in cols.items():
                values=[text(ws.cell(*anchor(ws,r,col)).value) for r in rows]
                rec[field]='\n'.join(dict.fromkeys(v for v in values if v))
            if 'Responsible' not in cols:
                legacy=[c.column for c in ws[4] if text(c.value).upper() in {'SPCL','SPOL','PM','SPECIAL OPS','PARTNER RESPONSIBLE','SPECIAL OPS RESPONSIBLE'}]
                rec['Responsible']='; '.join(dict.fromkeys(text(ws.cell(*anchor(ws,r,c)).value) for r in rows for c in legacy if text(ws.cell(*anchor(ws,r,c)).value)))
            rec.update(Partner=partner,SourceRow=start,Rows=rows)
            records.append(rec)
    keys=[(r['Partner'],r['Project'],r['Topic'],r['Problem']) for r in records]
    if len(keys)!=len(set(keys)):raise ValueError('Each issue must have a distinct partner, project, topic and problem.')
    return records

def save_issue(data,partner,start,fields):
    if partner not in PARTNERS or set(fields)!=set(FIELDS):raise ValueError('Unexpected issue fields.')
    if not text(fields['Project']) or not text(fields['Problem']):raise ValueError('Project and Problem are required.')
    records=read(data);wb=load_workbook(BytesIO(data));ws=wb[partner];cols=mapping(ws)
    old=next((x for x in records if x['Partner']==partner and x['SourceRow']==start),None) if start is not None else None
    if start is not None and old is None:raise ValueError('Issue no longer exists. Refresh first.')
    if old:rows=old['Rows']
    else:
        used=[c.row for row in ws for c in row if c.value is not None]
        start=max([4]+used)+1;rows=[start]
    changes=[]
    for field in FIELDS:
        new=text(fields[field]);previous=old[field] if old else ''
        if new==previous:continue
        if field not in cols:
            col=max([c.column for row in ws for c in row if c.value is not None]+[0])+1
            cols[field]=col;ws.cell(4,col,field)
            ws.cell(4,col)._style=copy(ws.cell(4,cols['Problem'])._style)
        col=cols[field]
        targets=list(dict.fromkeys(anchor(ws,r,col) for r in rows))
        for target in targets:
            for merged in ws.merged_cells.ranges:
                if (merged.min_row,merged.min_col)==target and (merged.min_row<min(rows) or merged.max_row>max(rows)):
                    raise ValueError(f'{field} is shared with other issues through merged Excel cells. Unmerge that field in Excel before changing it here.')
            if ws.cell(*target).data_type=='f':raise ValueError(f'{field} contains an Excel formula. Edit that formula in Excel.')
        for i,target in enumerate(targets):
            cell=ws.cell(*target)
            if i==0 and field in DATES and new:
                try:cell.value=date.fromisoformat(new)
                except ValueError as exc:raise ValueError(f'{field}: enter a valid date.') from exc
                cell.number_format='dd mmm yyyy'
            else:
                cell.value=new if i==0 else ''
                cell.data_type='s'
        changes.append((field,previous,new))
    if not changes:raise ValueError('No changes to save.')
    old_key=(partner,old['Project'],old['Topic']+' — '+old['Problem']) if old else None
    new_key=(partner,text(fields['Project']),text(fields['Topic'])+' — '+text(fields['Problem']))
    if old_key and old_key!=new_key:
        for sheet in ['Dashboard Actions','Dashboard History']:
            if sheet in wb:
                for row in wb[sheet].iter_rows(min_row=2):
                    if tuple(text(c.value) for c in row[:3])==old_key:
                        for c,v in zip(row[:3],new_key):c.value=v;c.data_type='s'
    history=wb['Dashboard History'] if 'Dashboard History' in wb else wb.create_sheet('Dashboard History')
    if history.cell(1,1).value is None:
        for c,v in enumerate(HISTORY,1):history.cell(1,c,v)
    if [c.value for c in history[1]][:len(HISTORY)]!=HISTORY:raise ValueError('Dashboard History has unexpected headers; cannot safely append history.')
    stamp=datetime.now(timezone.utc).strftime('%Y-%m-%dT%H:%M:%SZ')
    for field,previous,new in changes:
        values=list(new_key)+[stamp,'Anonymous editor',field,previous,new,'']
        r=history.max_row+1
        for c,v in enumerate(values,1):cell=history.cell(r,c,v);cell.data_type='s'
    out=BytesIO();wb.save(out);read(out.getvalue());return out.getvalue()
