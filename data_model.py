"""Workbook schema, validation and calculations. No issue ID field required."""
from io import BytesIO
import pandas as pd
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation

KEY = ['Partner', 'Project', 'Issue']
ISSUE_COLUMNS = KEY + ['Topic', 'Root cause', 'Root cause confidence', 'Problem', 'Effects', 'Solutions', 'Responsible', 'Severity', 'Status', 'Date identified', 'Target resolution', 'Decision needed', 'Decision owner', 'Decision due', 'Last updated', 'Success criterion', 'Resolution date', 'Pertinent notes']
ACTION_COLUMNS = KEY + ['Action', 'Responsible', 'Status', 'Due date', 'Completed date', 'Blocker', 'Pertinent notes']
UPDATE_COLUMNS = KEY + ['Date', 'Change type', 'Previous value', 'New value', 'Note']
SCHEMA = {'Issues': ISSUE_COLUMNS, 'Actions': ACTION_COLUMNS, 'Updates': UPDATE_COLUMNS}
DATES = {'Issues': ['Date identified','Target resolution','Decision due','Last updated','Resolution date'], 'Actions': ['Due date','Completed date'], 'Updates': ['Date']}
OPTIONS = {'Severity': ['Critical','High','Medium','Low','Unassessed'], 'Root cause confidence': ['Suspected','Confirmed'], 'Issues.Status': ['New','In Progress','Blocked','Monitoring','Closed','Unreviewed'], 'Actions.Status': ['Not Started','In Progress','Blocked','Done'], 'Change type': ['Created','Status','Severity','Owner','Deadline','Action completed','Reopened','Note']}


def demo(today):
    def date(days): return today + pd.Timedelta(days=days)
    issues = [
        ['Partner A','Expansion','Materials delayed','Supply chain','Supplier capacity','Confirmed','Materials have not arrived','Installation and launch at risk','Approve alternative supplier','Operations Manager','Critical','Blocked',date(-35),date(5),'Approve alternative quotation','Finance Manager',date(-2),date(-4),'Materials received and installation schedule confirmed',None,'Quotation is ready for review.'],
        ['Partner B','Training','Low attendance','Participation','Timing conflicts','Suspected','Attendance is below target','Training outcomes at risk','Pilot evening sessions','Programme Lead','High','In Progress',date(-25),date(14),'','',None,date(-17),'Attendance exceeds 80% for two sessions',None,'Confirm the cause with participants.'],
        ['Partner A','Expansion','Equipment reliability','Operations','Maintenance gap','Confirmed','Repeated equipment stoppages','Lost production time','Complete maintenance and monitor','Engineering Lead','Medium','Monitoring',date(-45),date(10),'','',None,date(-1),'14 consecutive days without an unplanned stoppage',None,'Maintenance completed; monitoring continues.'],
        ['Partner C','Reporting','Report approved','Reporting','Unclear review process','Confirmed','Report approval was delayed','Late reporting','Agree a review timetable','Project Lead','Low','Closed',date(-60),date(-10),'','',None,date(-8),'Partner accepts final report',date(-8),'Accepted by partner.'],
        ['Partner D','Onboarding','Missing onboarding plan','People','Unclear ownership','Suspected','No agreed onboarding plan','Start date at risk','Assign owner and publish plan','','High','New',date(-7),None,'Assign accountable owner','Director',date(2),None,'Owner assigned and plan agreed',None,''],
    ]
    actions = [
        ['Partner A','Expansion','Materials delayed','Approve alternative supplier','Finance Manager','Blocked',date(-2),None,'Budget approval',''],
        ['Partner A','Expansion','Materials delayed','Confirm revised delivery date','Operations Manager','Not Started',date(3),None,'Budget approval','Depends on supplier approval.'],
        ['Partner B','Training','Low attendance','Interview participants','Programme Lead','In Progress',date(-3),None,'',''],
        ['Partner A','Expansion','Equipment reliability','Complete maintenance','Engineering Lead','Done',date(-3),date(-2),'',''],
        ['Partner C','Reporting','Report approved','Obtain approval','Project Lead','Done',date(-10),date(-8),'',''],
    ]
    updates = [
        ['Partner A','Expansion','Materials delayed',date(-4),'Severity','High','Critical','Launch date is now at risk.'],
        ['Partner A','Expansion','Equipment reliability',date(-1),'Status','In Progress','Monitoring','Maintenance completed; checking effectiveness.'],
        ['Partner C','Reporting','Report approved',date(-8),'Status','In Progress','Closed','Report accepted.'],
    ]
    return {name: pd.DataFrame(rows, columns=cols) for name, rows, cols in [('Issues',issues,ISSUE_COLUMNS),('Actions',actions,ACTION_COLUMNS),('Updates',updates,UPDATE_COLUMNS)]}


def workbook_bytes(tables):
    output = BytesIO()
    with pd.ExcelWriter(output, engine='openpyxl', datetime_format='yyyy-mm-dd', date_format='yyyy-mm-dd') as writer:
        for name, cols in SCHEMA.items():
            tables.get(name, pd.DataFrame(columns=cols)).reindex(columns=cols).to_excel(writer, sheet_name=name, index=False)
            ws = writer.sheets[name]
            ws.freeze_panes = 'D2'
            ws.auto_filter.ref = ws.dimensions
            ws.row_dimensions[1].height = 32
            for cell in ws[1]:
                cell.font = Font(bold=True, color='FFFFFF')
                cell.fill = PatternFill('solid', fgColor='173D50')
                cell.alignment = Alignment(wrap_text=True)
                ws.column_dimensions[cell.column_letter].width = 24 if cell.value not in ['Problem','Effects','Solutions','Pertinent notes','Success criterion'] else 45
                opts = OPTIONS.get(f'{name}.{cell.value}', OPTIONS.get(cell.value))
                if opts:
                    dv = DataValidation(type='list', formula1='"'+','.join(opts)+'"', allow_blank=True)
                    dv.errorTitle = 'Choose a listed value'
                    dv.error = 'Use the dropdown options.'
                    dv.showErrorMessage = True
                    ws.add_data_validation(dv)
                    dv.add(f'{cell.column_letter}2:{cell.column_letter}5000')
            for row in ws.iter_rows(min_row=2):
                for cell in row:
                    cell.alignment = Alignment(vertical='top', wrap_text=True)
                    if isinstance(cell.value, str): cell.data_type = 's'
                    if cols[cell.column-1] in DATES[name]: cell.number_format = 'yyyy-mm-dd'
    return output.getvalue()


def read_workbook(data):
    return pd.read_excel(BytesIO(data), sheet_name=None, engine='openpyxl')


def validate(raw):
    result, errors = {}, []
    for name, cols in SCHEMA.items():
        if name not in raw and name == 'Issues': errors.append('Missing Issues sheet.')
        frame = raw.get(name, pd.DataFrame(columns=cols)).copy().dropna(how='all')
        frame.columns = [str(c).strip() for c in frame.columns]
        if frame.columns.duplicated().any():
            errors.append(f'{name}: duplicate column names.'); continue
        missing = [c for c in cols if c not in frame]
        if missing: errors.append(f'{name}: missing columns: {", ".join(missing)}.')
        frame = frame.reindex(columns=cols)
        for c in cols:
            if c in DATES[name]:
                original = frame[c]
                converted = pd.to_datetime(original, errors='coerce')
                bad = original.notna() & original.astype(str).str.strip().ne('') & converted.isna()
                if bad.any(): errors.append(f'{name}: invalid {c} in Excel rows {list(frame.index[bad] + 2)}. Use Excel dates or YYYY-MM-DD.')
                frame[c] = converted
            else:
                frame[c] = frame[c].fillna('').astype(str).str.strip()
                options = OPTIONS.get(f'{name}.{c}', OPTIONS.get(c))
                if options:
                    mapping = {v.casefold():v for v in options}
                    frame[c] = frame[c].map(lambda v: mapping.get(v.casefold(), v))
                    bad = frame[c].ne('') & ~frame[c].isin(options)
                    if bad.any(): errors.append(f'{name}: invalid {c}: {", ".join(frame.loc[bad,c].unique())}.')
        for c in KEY:
            if frame[c].eq('').any(): errors.append(f'{name}: every row needs {c}.')
        for c in ({'Issues':['Severity','Status'], 'Actions':['Action','Status'], 'Updates':['Change type']}[name]):
            if frame[c].eq('').any(): errors.append(f'{name}: every row needs {c}.')
        if name == 'Updates' and frame['Date'].isna().any(): errors.append('Updates: every row needs Date.')
        result[name] = frame
    if errors: raise ValueError('\n'.join(errors))
    issues = result['Issues']
    if issues.duplicated(KEY).any(): errors.append('Issues: Partner + Project + Issue must be distinct. Use descriptive issue names to distinguish problems.')
    keys = set(map(tuple, issues[KEY].values))
    for name in ['Actions','Updates']:
        unknown = [i+2 for i,row in result[name].iterrows() if tuple(row[KEY]) not in keys]
        if unknown: errors.append(f'{name}: no matching issue for Excel rows {unknown}. Match Partner, Project and Issue exactly.')
    if errors: raise ValueError('\n'.join(errors))
    return result


def calculate(tables, today, stale_days=14):
    issues, actions = tables['Issues'].copy(), tables['Actions'].copy()
    actions['Overdue'] = actions['Status'].ne('Done') & actions['Due date'].lt(today)
    actions['Days overdue'] = (today-actions['Due date']).dt.days.clip(lower=0).fillna(0).astype(int)
    issues['Open'] = issues['Status'].ne('Closed')
    issues['Age (days)'] = (today-issues['Date identified']).dt.days.clip(lower=0)
    issues['Impact level'] = issues['Severity'].map({'Low':1,'Medium':2,'High':3,'Critical':4})
    issues['Decision pending'] = issues['Open'] & issues['Decision needed'].ne('')
    issues['Resolution overdue'] = issues['Open'] & issues['Target resolution'].lt(today)
    issues['Stale'] = issues['Open'] & (issues['Last updated'].isna() | (today-issues['Last updated']).dt.days.gt(stale_days))
    counts = actions.groupby(KEY)['Overdue'].sum().rename('Overdue actions')
    pending = actions[actions['Status'].ne('Done')].groupby(KEY).size().rename('Pending actions')
    issues = issues.join(counts, on=KEY).join(pending,on=KEY)
    issues[['Overdue actions','Pending actions']] = issues[['Overdue actions','Pending actions']].fillna(0).astype(int)
    def quality(row):
        flags = []
        if row['Severity'] == 'Unassessed': flags.append('Severity not assessed')
        if row['Status'] == 'Unreviewed': flags.append('Status not reviewed')
        if row['Open']:
            for c in ['Responsible','Success criterion']:
                if not row[c]: flags.append('Missing '+c.lower())
            for c in ['Date identified','Target resolution','Last updated']:
                if pd.isna(row[c]): flags.append('Missing '+c.lower())
            if row['Status'] != 'Monitoring' and row['Pending actions'] == 0: flags.append('No pending action')
            if row['Stale']: flags.append('Update needed')
        if row['Decision pending'] and (not row['Decision owner'] or pd.isna(row['Decision due'])): flags.append('Decision owner/date missing')
        if row['Status'] == 'Closed' and pd.isna(row['Resolution date']): flags.append('Closed without resolution date')
        if row['Status'] == 'Closed' and row['Pending actions'] > 0: flags.append('Closed with pending actions')
        if row['Open'] and pd.notna(row['Resolution date']): flags.append('Open with resolution date')
        if pd.notna(row['Resolution date']) and pd.notna(row['Date identified']) and row['Resolution date'] < row['Date identified']: flags.append('Resolution precedes identification')
        if pd.notna(row['Date identified']) and row['Date identified'] > today: flags.append('Identification date is in the future')
        return '; '.join(flags)
    issues['Data gaps'] = issues.apply(quality,axis=1) if len(issues) else pd.Series(dtype=str)
    def reasons(row):
        flags = [row['Severity']+' impact']
        if row['Overdue actions']: flags.append(f"{row['Overdue actions']} overdue action(s)")
        if row['Resolution overdue']: flags.append('Resolution overdue')
        if row['Status']=='Blocked': flags.append('Blocked')
        if row['Decision pending']: flags.append('Decision needed')
        if row['Stale']: flags.append('Update needed')
        return ' · '.join(flags)
    issues['Attention reasons'] = issues.apply(reasons,axis=1) if len(issues) else pd.Series(dtype=str)
    return issues, actions
