"""Offline checks; uses synthetic workbook data and mocked Dropbox responses."""
import json
import unittest
from datetime import datetime,timezone
from io import BytesIO
from unittest.mock import patch
import requests
import pandas as pd
from openpyxl import Workbook,load_workbook
from workbook_model import read,patch_issue,add_action,add_issue,metadata_columns
from workbook_store import DropboxStore,ConflictError,UnknownSaveError,content_hash
from auth import hash_password,verify_password

NOW=datetime(2026,10,3,12,tzinfo=timezone.utc)


def fixture():
    w=Workbook();w.active.title='Overview';w['Overview']['A1']='=SPCL!D5'
    for name in ['SPCL','SPOL','PM']:
        s=w.create_sheet(name)
        for i,h in enumerate(['PROJECT','TOPIC','ROOT CAUSE','PROBLEM','EFFECTS','SOLUTION','SPCL','SPECIAL OPS','TIMING','NOTES'],1): s.cell(4,i,h)
    s=w['SPCL']
    for i,v in enumerate(['Farm','Equipment','Maintenance gap','Equipment stops','Delayed work','Inspect equipment','Team A','Team B','Next week','Keep this note'],1):s.cell(5,i,v)
    for col in ['A','B','C','D','J']: s.merge_cells(f'{col}5:{col}7')
    s['E6']='Additional effect';s['E7']='Third effect';s['G6']='Other owner'
    s['K2']='Preserve unrelated data'
    b=BytesIO();w.save(b);return b.getvalue()


def change(data,fields=None,details=None,extras=None):
    p=read(data);g=p['groups'][('SPCL',5)]
    return patch_issue(data,'SPCL',5,fields or {'Severity':'High'},details or g['details'],extras or [],'tester',NOW)


class ModelTests(unittest.TestCase):
    def setUp(self): self.data=fixture()
    def test_merged_issue(self):
        p=read(self.data)
        self.assertEqual(len(p['tables']['Issues']),1)
        self.assertEqual(len(p['groups'][('SPCL',5)]['details']),3)
        self.assertIn('Third effect',p['tables']['Issues'].iloc[0]['Effects'])
        self.assertEqual(p['tables']['Issues'].iloc[0]['Severity'],'Unassessed')
    def test_preserve_source(self):
        changed=change(self.data)
        before=load_workbook(BytesIO(self.data));after=load_workbook(BytesIO(changed))
        for sheet in ['Overview','SPCL','SPOL','PM']:
            self.assertEqual(set(map(str,before[sheet].merged_cells.ranges)),set(map(str,after[sheet].merged_cells.ranges)))
            for row in before[sheet]:
                for cell in row:
                    self.assertEqual(cell.value,after[sheet][cell.coordinate].value)
                    self.assertEqual(cell._style,after[sheet][cell.coordinate]._style)
        self.assertEqual(after['Overview']['A1'].data_type,'f')
        self.assertIn('Dashboard History',after.sheetnames)
    def test_detail_update_and_literal_formula(self):
        p=read(self.data);d=p['groups'][('SPCL',5)]['details'];d[1]['Effects']='=literal text'
        out=change(self.data,details=d);w=load_workbook(BytesIO(out))
        self.assertEqual(w['SPCL']['E6'].value,'=literal text');self.assertEqual(w['SPCL']['E6'].data_type,'s')
        self.assertEqual(w['SPCL']['E5'].value,'Delayed work')
    def test_add_action_rename_and_edit(self):
        out=add_action(self.data,'SPCL',5,{'Action':'Order parts','Responsible':'Owner','Status':'Not Started'},'tester',NOW)
        out=change(out,{'Project':'New farm','Problem':'Revised description'})
        p=read(out);self.assertEqual(p['extra'].iloc[0]['Project'],'New farm')
        row=p['extra'].iloc[0].to_dict();row['Excel row']=p['extra'].attrs['excel_rows'][0];row['Status']='Done';row['Completed date']=NOW.date()
        out=change(out,{'Severity':'Critical'},extras=[row])
        self.assertEqual(read(out)['extra'].iloc[0]['Status'],'Done')
        self.assertTrue(read(out)['history']['Project'].eq('New farm').all())
    def test_add_issue_empty_partner(self):
        out=add_issue(self.data,'SPOL',{'Project':'Launch','Problem':'Permit pending','Status':'New','Severity':'High'},{'Effects':'Delay'},'tester',NOW)
        p=read(out);self.assertEqual(len(p['tables']['Issues']),2)
        self.assertEqual(p['tables']['Issues'].iloc[-1]['Partner'],'SPOL')
    def test_close_requires_completed_work(self):
        with self.assertRaisesRegex(ValueError,'pending actions'):change(self.data,{'Status':'Closed','Success criterion':'Works','Resolution date':NOW.date()})
        d=read(self.data)['groups'][('SPCL',5)]['details'];d[0]['Action status']='Done';d[0]['Action completed']=NOW.date()
        out=change(self.data,{'Status':'Closed','Success criterion':'Works','Resolution date':NOW.date()},d)
        self.assertEqual(read(out)['tables']['Issues'].iloc[0]['Status'],'Closed')
        out=change(out,{'Status':'In Progress','Resolution date':None})
        self.assertIn('Reopened',read(out)['history']['Field'].tolist())
    def test_duplicate_issue_rejected(self):
        with self.assertRaisesRegex(ValueError,'Duplicate'):
            add_issue(self.data,'SPCL',{'Project':'Farm','Topic':'Equipment','Problem':'Equipment stops'},{},'tester',NOW)
    def test_invalid_date_rejected(self):
        with self.assertRaises(ValueError):change(self.data,{'Target resolution':'next week'})


class Response:
    def __init__(self,status=200,data=None): self.status_code=status;self.data=data or {}
    def json(self):return self.data


class FakeHTTP:
    def __init__(self,responses):self.responses=iter(responses);self.calls=[]
    def post(self,url,**kwargs):
        self.calls.append((url,kwargs));r=next(self.responses)
        if isinstance(r,Exception):raise r
        return r


class StoreTests(unittest.TestCase):
    def setUp(self):self.snapshot={'file_id':'id:test','rev':'old','bytes':b'old'}
    def store(self,responses):
        http=FakeHTTP(responses);s=DropboxStore({},http);s.token='fake';return s,http
    def test_revision_precheck_prevents_write(self):
        s,h=self.store([Response(data={'rev':'new'})])
        with self.assertRaises(ConflictError):s.save(self.snapshot,b'new')
        self.assertEqual(len(h.calls),1)
    def test_atomic_conflict_after_precheck(self):
        s,h=self.store([Response(data={'rev':'old'}),Response(409)])
        with self.assertRaises(ConflictError):s.save(self.snapshot,b'new')
        args=json.loads(h.calls[1][1]['headers']['Dropbox-API-Arg'])
        self.assertEqual(args['mode'],{'.tag':'update','update':'old'})
        self.assertTrue(args['strict_conflict']);self.assertFalse(args['autorename'])
    def test_success(self):
        s,h=self.store([Response(data={'rev':'old'}),Response(data={'rev':'new'})])
        self.assertEqual(s.save(self.snapshot,b'new')['rev'],'new')
    def test_timeout_confirmed_by_content_hash(self):
        s,h=self.store([Response(data={'rev':'old'}),requests.Timeout(),Response(data={'rev':'new','content_hash':content_hash(b'new')})])
        self.assertEqual(s.save(self.snapshot,b'new')['rev'],'new')
    def test_uncertain_timeout_does_not_retry(self):
        s,h=self.store([Response(data={'rev':'old'}),requests.Timeout(),Response(data={'rev':'old','content_hash':'other'})])
        with self.assertRaises(UnknownSaveError):s.save(self.snapshot,b'new')
        self.assertEqual(sum('files/upload' in u for u,k in h.calls),1)
    def test_passwords(self):
        encoded=hash_password('a long test password')
        self.assertTrue(verify_password('a long test password',encoded));self.assertFalse(verify_password('wrong',encoded))

if __name__=='__main__':unittest.main()
