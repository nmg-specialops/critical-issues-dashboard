"""Read and patch the user's partner-sheet layout without rebuilding the workbook."""
from copy import copy
from datetime import datetime, date
from io import BytesIO
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.workbook.properties import CalcProperties
from openpyxl.utils import get_column_letter
from data_model import ISSUE_COLUMNS, ACTION_COLUMNS, UPDATE_COLUMNS

PARTNERS=['SPCL','SPOL','PM']
CORE={'Project':1,'Topic':2,'Root cause':3,'Problem':4,'Effects':5,'Solutions':6,'Partner responsible':7,'Special Ops responsible':8,'Timing':9,'Pertinent notes':10}
META=['Severity','Status','Responsible','Date identified','Target resolution','Decision needed','Decision owner','Decision due','Last updated','Success criterion','Resolution date','Root cause confidence']
ACTION_META=['Action owner','Action status','Action due','Action completed','Action blocker']
DATE_FIELDS={'Date identified','Target resolution','Decision due','Last updated','Resolution date','Action due','Action completed','Due date','Completed date'}
ISSUE_STATUSES=['Unreviewed','New','In Progress','Blocked','Monitoring','Closed']
SEVERITIES=['Unassessed','Critical','High','Medium','Low']
ACTION_STATUSES=['Unreviewed','Not Started','In Progress','Blocked','Done']
NATURAL=['Partner','Project','Issue']
EXTRA_SHEET='Dashboard Actions'
HISTORY_SHEET='Dashboard History'
EXTRA_COLS=NATURAL+['Action','Responsible','Status','Due date','Completed date','Blocker','Pertinent notes']
HISTORY_COLS=NATURAL+['Timestamp UTC','User','Field','Previous value','New value','Note']
DETAIL_COLS=['Excel row','Effects','Solutions','Partner responsible','Special Ops responsible','Timing']+ACTION_META


def text(value):
    if value is None or (not isinstance(value,(list,dict)) and pd.isna(value)): return ''
    if isinstance(value,(datetime,date,pd.Timestamp)): return value.isoformat()
    return str(value).strip()


def parse_date(value, label='Date'):
    if not text(value): return None
    if isinstance(value,(int,float)): raise ValueError(f'{label}: use an Excel date, not an unformatted number.')
    try:
        v=pd.Timestamp(value)
        if v.tzinfo: v=v.tz_convert(None)
        return v.normalize().to_pydatetime()
    except Exception as exc: raise ValueError(f'{label}: use an Excel date or YYYY-MM-DD.') from exc


def normalized(value,field):
    return parse_date(value,field) if field in DATE_FIELDS else text(value)


def metadata_columns(ws,create=False):
    found={}
    for cell in ws[4]:
        if text(cell.value) in META+ACTION_META:
            if text(cell.value) in found: raise ValueError(f'{ws.title}: duplicate tracking header {cell.value}.')
            found[text(cell.value)]=cell.column
    if create:
        used=[c.column for row in ws for c in row if c.value is not None]
        next_col=max([10]+used)+1
        for name in META+ACTION_META:
            if name not in found:
                found[name]=next_col
                cell=ws.cell(4,next_col,name)
                cell.font=Font(bold=True,color='FFFFFF')
                cell.fill=PatternFill('solid',fgColor='173D50')
                cell.alignment=Alignment(wrap_text=True,vertical='top')
                ws.column_dimensions[get_column_letter(next_col)].width=24
                next_col+=1
        ws.row_dimensions[4].height=max(ws.row_dimensions[4].height or 15,32)
    return found


def anchor(ws,row,col):
    for merged in ws.merged_cells.ranges:
        if merged.min_row<=row<=merged.max_row and merged.min_col<=col<=merged.max_col:
            return merged.min_row,merged.min_col
    return row,col


def value(ws,row,col):
    r,c=anchor(ws,row,col)
    return ws.cell(r,c).value


def issue_label(topic,problem): return f'{text(topic)} — {text(problem)}'


def auxiliary(wb,name,columns):
    if name not in wb: return pd.DataFrame(columns=columns)
    ws=wb[name]
    if [ws.cell(1,c+1).value for c in range(len(columns))]!=columns:
        raise ValueError(f'{name}: unexpected headers. Restore the dashboard-created header row before saving.')
    records=[]; excel_rows=[]
    for row in ws.iter_rows(min_row=2,max_col=len(columns)):
        if any(c.value is not None for c in row):
            records.append([c.value for c in row]); excel_rows.append(row[0].row)
    frame=pd.DataFrame(records,columns=columns).fillna('')
    frame.attrs['excel_rows']=excel_rows
    return frame


def read(data):
    wb=load_workbook(BytesIO(data))
    issues=[]; actions=[]; groups={}
    for partner in PARTNERS:
        if partner not in wb: raise ValueError(f'Missing partner sheet: {partner}.')
        ws=wb[partner]
        expected=['PROJECT','TOPIC','ROOT CAUSE','PROBLEM','EFFECTS','SOLUTION','SPCL','SPECIAL OPS','TIMING','NOTES']
        actual=[text(ws.cell(4,c).value).upper() for c in range(1,11)]
        # The partner-responsible header may be relabelled for SPOL/PM.
        if actual[:6]!=expected[:6] or actual[7:]!=expected[7:]:
            raise ValueError(f'{partner}: expected the original ten-column layout with headers in row 4.')
        cols=metadata_columns(ws)
        sheet_groups={}
        for r in range(5,ws.max_row+1):
            if not any(ws.cell(r,c).value is not None for c in range(1,11)): continue
            start,_=anchor(ws,r,4)
            if not text(value(ws,r,4)):
                raise ValueError(f'{partner} row {r}: enter a Problem or merge it into the relevant issue block.')
            sheet_groups.setdefault(start,[]).append(r)
        for start,physical in sheet_groups.items():
            # Include all rows of a merged issue block, even if some details are empty.
            group_end=start
            for merged in ws.merged_cells.ranges:
                if merged.min_col<=4<=merged.max_col and merged.min_row==start: group_end=max(group_end,merged.max_row)
            rows=list(range(start,max(group_end,max(physical))+1))
            rec={c:'' for c in ISSUE_COLUMNS}
            rec.update(Partner=partner,Project=text(value(ws,start,1)),Topic=text(value(ws,start,2)),Problem=text(value(ws,start,4)))
            if not rec['Project']: raise ValueError(f'{partner} row {start}: missing Project.')
            rec['Issue']=issue_label(rec['Topic'],rec['Problem'])
            for c in ['Root cause','Pertinent notes']: rec[c]=text(value(ws,start,CORE[c]))
            for c in ['Effects','Solutions']:
                vals=[text(value(ws,r,CORE[c])) for r in rows]
                rec[c]='\n'.join(dict.fromkeys(v for v in vals if v))
            for c in META:
                v=ws.cell(start,cols[c]).value if c in cols else None
                rec[c]=normalized(v,c)
            rec['Severity']=rec['Severity'] or 'Unassessed'
            rec['Status']=rec['Status'] or 'Unreviewed'
            if rec['Severity'] not in SEVERITIES: raise ValueError(f'{partner} row {start}: invalid Severity.')
            if rec['Status'] not in ISSUE_STATUSES: raise ValueError(f'{partner} row {start}: invalid Status.')
            details=[]
            for r in rows:
                d={'Excel row':r}
                for c in ['Effects','Solutions','Partner responsible','Special Ops responsible','Timing']: d[c]=text(value(ws,r,CORE[c]))
                for c in ACTION_META:
                    v=ws.cell(r,cols[c]).value if c in cols else None
                    d[c]=normalized(v,c)
                d['Action status']=d['Action status'] or 'Unreviewed'
                if d['Action status'] not in ACTION_STATUSES: raise ValueError(f'{partner} row {r}: invalid Action status.')
                details.append(d)
                if d['Solutions']:
                    a={c:'' for c in ACTION_COLUMNS}
                    a.update({c:rec[c] for c in NATURAL})
                    a.update(Action=d['Solutions'],Responsible=d['Action owner'],Status=d['Action status'],**{'Due date':d['Action due'],'Completed date':d['Action completed'],'Blocker':d['Action blocker']})
                    a['Pertinent notes']=f"Partner: {d['Partner responsible']} | Special Ops: {d['Special Ops responsible']} | Timing: {d['Timing']}"
                    actions.append(a)
            rec['Source row']=start
            rec['Partner responsible']='; '.join(dict.fromkeys(d['Partner responsible'] for d in details if d['Partner responsible']))
            rec['Special Ops responsible']='; '.join(dict.fromkeys(d['Special Ops responsible'] for d in details if d['Special Ops responsible']))
            rec['Timing']='; '.join(dict.fromkeys(d['Timing'] for d in details if d['Timing']))
            issues.append(rec)
            groups[(partner,start)]={'issue':rec,'details':details,'rows':rows}
    issue_df=pd.DataFrame(issues,columns=ISSUE_COLUMNS+['Source row','Partner responsible','Special Ops responsible','Timing'])
    if issue_df.duplicated(NATURAL).any(): raise ValueError('Partner + Project + Topic + Problem must distinguish each issue. Duplicate issue descriptions cannot be linked reliably without an ID.')
    extra=auxiliary(wb,EXTRA_SHEET,EXTRA_COLS)
    valid_keys=set(map(tuple,issue_df[NATURAL].values))
    for idx,row in extra.iterrows():
        if tuple(text(row[c]) for c in NATURAL) not in valid_keys: raise ValueError(f'{EXTRA_SHEET} row {idx+2}: no matching partner/project/issue. Restore the matching labels or rename the issue through the app.')
        a={c:normalized(row[c],c) for c in EXTRA_COLS}
        if a['Status'] not in ACTION_STATUSES: raise ValueError(f'{EXTRA_SHEET} row {idx+2}: invalid action status.')
        if not a['Action']: raise ValueError(f'{EXTRA_SHEET} row {idx+2}: missing Action.')
        actions.append(a)
    history=auxiliary(wb,HISTORY_SHEET,HISTORY_COLS)
    updates=[]
    for _,row in history.iterrows():
        kind={'Status':'Status','Severity':'Severity','Responsible':'Owner','Target resolution':'Deadline','Created':'Created','Reopened':'Reopened'}.get(row['Field'],'Note')
        if ('Action status' in row['Field'] or row['Field'].endswith(': Status')) and row['New value']=='Done': kind='Action completed'
        updates.append({**{c:text(row[c]) for c in NATURAL},'Date':parse_date(row['Timestamp UTC'],'History timestamp'),'Change type':kind,'Previous value':text(row['Previous value']),'New value':text(row['New value']),'Note':f"{row['User']} · {row['Field']}: {text(row['Note'])}"})
    action_df=pd.DataFrame(actions,columns=ACTION_COLUMNS)
    update_df=pd.DataFrame(updates,columns=UPDATE_COLUMNS)
    for frame,dates in [(issue_df,[c for c in META if c in DATE_FIELDS]),(action_df,['Due date','Completed date']),(update_df,['Date'])]:
        for c in dates: frame[c]=pd.to_datetime(frame[c],errors='raise')
    return {'workbook':wb,'groups':groups,'extra':extra,'history':history,'tables':{'Issues':issue_df,'Actions':action_df,'Updates':update_df}}


def get_or_create(wb,name,columns):
    if name in wb:
        auxiliary(wb,name,columns)
        return wb[name]
    ws=wb.create_sheet(name)
    ws.append(columns)
    for cell in ws[1]:
        cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='173D50')
        ws.column_dimensions[cell.column_letter].width=28
    ws.freeze_panes='D2'
    return ws


def set_literal(cell,new):
    cell.value=new if new is not None and new!='' else None
    if isinstance(new,str) and new!='': cell.data_type='s'
    cell.alignment=Alignment(wrap_text=True,vertical='top')
    if isinstance(new,(date,datetime)): cell.number_format='yyyy-mm-dd'


def assign(ws,r,c,new,changes,label):
    rr,cc=anchor(ws,r,c)
    old=ws.cell(rr,cc).value
    if text(old)==text(new): return
    if ws.cell(rr,cc).data_type=='f': raise ValueError(f'{ws.title}!{ws.cell(rr,cc).coordinate} contains a formula. Edit the source of that formula in Excel.')
    set_literal(ws.cell(rr,cc),new)
    changes.append((label,text(old),text(new)))


def log(wb,key,changes,user,now,note=''):
    ws=get_or_create(wb,HISTORY_SHEET,HISTORY_COLS)
    for field,old,new in changes:
        values=list(key)+[now.strftime('%Y-%m-%dT%H:%M:%SZ'),user,field,old,new,note]
        r=ws.max_row+1
        for c,v in enumerate(values,1): set_literal(ws.cell(r,c),v)
    ws.auto_filter.ref=ws.dimensions


def rekey(wb,old_key,new_key):
    if old_key==new_key: return
    for name,columns in [(EXTRA_SHEET,EXTRA_COLS),(HISTORY_SHEET,HISTORY_COLS)]:
        if name not in wb: continue
        ws=get_or_create(wb,name,columns)
        for r in range(2,ws.max_row+1):
            if tuple(text(ws.cell(r,c).value) for c in range(1,4))==old_key:
                for c,v in enumerate(new_key,1): set_literal(ws.cell(r,c),v)


def export_and_validate(wb,partner,start):
    wb.calculation=CalcProperties(calcId=191029,fullCalcOnLoad=True,forceFullCalc=True)
    output=BytesIO();wb.save(output);data=output.getvalue()
    parsed=read(data)
    rec=parsed['groups'][(partner,start)]['issue']
    if rec['Status']=='Closed':
        key=tuple(rec[c] for c in NATURAL)
        actions=parsed['tables']['Actions']
        pending=actions[actions[NATURAL].apply(lambda row: tuple(row)==key,axis=1) & actions['Status'].ne('Done')]
        if len(pending): raise ValueError('Complete the pending actions before closing the issue.')
        if not rec['Success criterion'] or not rec['Resolution date']: raise ValueError('Closing requires a Success criterion and Resolution date.')
    if rec['Status']!='Closed' and rec['Resolution date']: raise ValueError('Clear Resolution date when reopening an issue.')
    if rec['Date identified'] and rec['Resolution date'] and rec['Resolution date']<rec['Date identified']: raise ValueError('Resolution date cannot precede Date identified.')
    return data


def patch_issue(data,partner,start,fields,details,extra_edits,user,now,note=''):
    parsed=read(data);wb=parsed['workbook'];ws=wb[partner]
    group=parsed['groups'][(partner,start)];old=group['issue']
    old_key=tuple(old[c] for c in NATURAL)
    cols=metadata_columns(ws,create=True);changes=[]
    for field,new in fields.items():
        if field not in ['Project','Topic','Root cause','Problem','Pertinent notes']+META: raise ValueError('Unexpected issue field.')
        assign(ws,start,CORE[field] if field in CORE else cols[field],normalized(new,field),changes,field)
    if not text(ws.cell(start,1).value) or not text(ws.cell(start,4).value): raise ValueError('Project and Problem cannot be blank.')
    if len(details)!=len(group['details']): raise ValueError('Keep existing detail rows. Use Add action for additional work.')
    touched={}
    for item in details:
        r=int(item['Excel row'])
        if r not in group['rows']: raise ValueError('Invalid detail row.')
        for field in DETAIL_COLS[1:]:
            col=CORE[field] if field in CORE else cols[field]
            new=normalized(item.get(field),field)
            target=anchor(ws,r,col)
            if target in touched and text(touched[target])!=text(new): raise ValueError(f'{field} is merged across rows in Excel. Keep its values identical or edit the merge in Excel.')
            touched[target]=new
            assign(ws,r,col,new,changes,f'Row {r}: {field}')
    extra_ws=get_or_create(wb,EXTRA_SHEET,EXTRA_COLS) if extra_edits else None
    for edit in extra_edits:
        r=int(edit['Excel row'])
        if r<2 or r>extra_ws.max_row or tuple(text(extra_ws.cell(r,c).value) for c in range(1,4))!=old_key: raise ValueError('Additional action no longer matches the selected issue.')
        for field in EXTRA_COLS[3:]:
            assign(extra_ws,r,EXTRA_COLS.index(field)+1,normalized(edit.get(field),field),changes,f'Additional action row {r}: {field}')
    if not changes: raise ValueError('No changes to save.')
    assign(ws,start,cols['Last updated'],now.replace(tzinfo=None),changes,'Last updated')
    new_key=(partner,text(ws.cell(start,1).value),issue_label(ws.cell(start,2).value,ws.cell(start,4).value))
    rekey(wb,old_key,new_key)
    if old['Status']=='Closed' and fields.get('Status',old['Status'])!='Closed': changes.append(('Reopened','Closed',fields['Status']))
    log(wb,new_key,changes,user,now,note)
    return export_and_validate(wb,partner,start)


def add_action(data,partner,start,fields,user,now):
    parsed=read(data);wb=parsed['workbook'];rec=parsed['groups'][(partner,start)]['issue']
    if rec['Status']=='Closed': raise ValueError('Reopen this issue before adding an action.')
    if not text(fields.get('Action')): raise ValueError('Enter an action.')
    key=tuple(rec[c] for c in NATURAL)
    ws=get_or_create(wb,EXTRA_SHEET,EXTRA_COLS);r=ws.max_row+1
    values={**dict(zip(NATURAL,key)),**fields}
    for c,field in enumerate(EXTRA_COLS,1): set_literal(ws.cell(r,c),normalized(values.get(field,''),field))
    ws.auto_filter.ref=ws.dimensions
    pws=wb[partner];cols=metadata_columns(pws,True)
    set_literal(pws.cell(start,cols['Last updated']),now.replace(tzinfo=None))
    log(wb,key,[('Action added','',text(fields['Action']))],user,now)
    return export_and_validate(wb,partner,start)


def add_issue(data,partner,fields,detail,user,now):
    if partner not in PARTNERS: raise ValueError('Choose a known partner.')
    if not text(fields.get('Project')) or not text(fields.get('Problem')): raise ValueError('Project and Problem are required.')
    parsed=read(data);wb=parsed['workbook'];ws=wb[partner];cols=metadata_columns(ws,True)
    # Append after the existing formatted/merged area; no row insertion that could break formulas.
    r=max(ws.max_row,max([4]+[m.max_row for m in ws.merged_cells.ranges]))+1
    for c in range(1,ws.max_column+1):
        source=ws.cell(5,c);target=ws.cell(r,c)
        if source.has_style: target._style=copy(source._style)
    for field,new in fields.items():
        if field not in ['Project','Topic','Root cause','Problem','Pertinent notes']+META: raise ValueError('Unexpected issue field.')
        set_literal(ws.cell(r,CORE[field] if field in CORE else cols[field]),normalized(new,field))
    for field,new in detail.items():
        if field not in ['Effects','Solutions','Partner responsible','Special Ops responsible','Timing']+ACTION_META: raise ValueError('Unexpected detail field.')
        set_literal(ws.cell(r,CORE[field] if field in CORE else cols[field]),normalized(new,field))
    set_literal(ws.cell(r,cols['Last updated']),now.replace(tzinfo=None))
    key=(partner,text(fields['Project']),issue_label(fields.get('Topic',''),fields['Problem']))
    log(wb,key,[('Created','',fields['Problem'])],user,now)
    return export_and_validate(wb,partner,r)
